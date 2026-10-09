#!/usr/bin/env python3
"""Создаёт обезличенные учебные выгрузки для Excel и Google Таблиц."""
from __future__ import annotations

import csv
import math
import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "downloads"
RNG = random.Random(13026)

CHANNELS = {
    "Портал": ["Портал", " портал ", "ПОРТАЛ", "электронный портал"],
    "Контакт-центр": ["Контакт-центр", "контакт центр", "КЦ", " горячая линия "],
    "Письменное обращение": ["Письменное обращение", "письмо", "письменно", "ПОЧТА"],
    "Личный приём": ["Личный приём", "личный прием", "приём", "ОЧНО"],
}
CATEGORIES = ["Благоустройство", "Транспортная доступность", "Социальная поддержка", "Иное"]
TEXTS = [
    "Запрос разъяснения порядка предоставления услуги",
    "Сообщение о состоянии объекта общего пользования",
    "Предложение об изменении маршрута обслуживания",
    "Запрос сведений о доступных мерах поддержки",
    "Уточнение статуса ранее направленного обращения",
]


def fmt_date(value: date | None, row: int) -> str:
    if value is None:
        return "н/д" if row % 3 else ""
    return value.strftime("%d.%m.%Y") if row % 2 else value.isoformat()


def make_rows() -> list[list[object]]:
    rows: list[list[object]] = []
    start = date(2026, 1, 1)
    for i in range(1, 12001):
        registered = start + timedelta(days=RNG.randrange(0, 240))
        target = RNG.choice([5, 10, 15, 30])
        closed = RNG.random() > 0.13
        duration = max(0, int(RNG.gauss(target * 0.82, max(2, target * 0.38))))
        completed = registered + timedelta(days=duration) if closed else None
        channel = RNG.choice(list(CHANNELS))
        status = "Завершено" if closed else RNG.choice(["В работе", "Требуется уточнение"])
        row = [
            f"EDU-{i:06d}",
            fmt_date(registered, i),
            fmt_date(completed, i + 1),
            RNG.choice(CHANNELS[channel]),
            RNG.choice(CATEGORIES),
            status,
            target,
            f"Территориальная группа {RNG.choice('ABCDEF')}",
            RNG.choice(TEXTS),
        ]
        rows.append(row)
        if i % 37 == 0:
            rows.append(row.copy())
    RNG.shuffle(rows)
    return rows


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    headers = ["ID обращения", "Дата регистрации", "Дата завершения", "Канал", "Категория", "Статус", "Целевой срок, дней", "Территориальная группа", "Краткое содержание"]
    rows = make_rows()

    csv_path = OUT / "obrashcheniya_12000_google.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh, delimiter=";")
        writer.writerow(headers)
        writer.writerows(rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "Выгрузка"
    ws.append(headers)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = [17, 19, 19, 23, 27, 23, 21, 26, 54]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + idx)].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173A63")
        cell.alignment = Alignment(wrap_text=True)
    table = Table(displayName="TrainingAppeals", ref=ws.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)

    info = wb.create_sheet("Описание")
    notes = [
        ("Назначение", "Обезличенная учебная выгрузка для отработки автоматизации в Excel."),
        ("Объём", f"{len(rows)} строк с преднамеренными полными повторами."),
        ("Особенности", "Даты записаны в двух форматах; каналы имеют варианты написания; незавершённые обращения содержат пустую дату или «н/д»."),
        ("Ограничение", "Данные не описывают фактическую работу организаций, территорий или должностных лиц."),
    ]
    for row in notes:
        info.append(row)
    info.column_dimensions["A"].width = 20
    info.column_dimensions["B"].width = 105
    info["A1"].font = Font(bold=True)
    wb.save(OUT / "obrashcheniya_12000_excel.xlsx")

    with (OUT / "seasonality_training.csv").open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh, delimiter=";")
        writer.writerow(["дата", "показатель", "базовый_уровень", "долгосрочное_изменение", "годовая_волна", "короткая_волна"])
        start_day = date(2024, 1, 1)
        wave_rng = random.Random(27)
        for t in range(1095):
            level = 100
            long_change = -5 + 15 * t / 1094
            annual = 17 * math.sin(2 * math.pi * (t - 55) / 365)
            short_wave = 5 * math.sin(2 * math.pi * t / 30.4)
            value = level + long_change + annual + short_wave + (wave_rng.random() - .5) * 3.4
            writer.writerow([(start_day + timedelta(days=t)).isoformat(), round(value, 2), level, round(long_change, 2), round(annual, 2), round(short_wave, 2)])

    with (OUT / "experiment_training.csv").open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh, delimiter=";")
        writer.writerow(["момент_наблюдения", "вариант", "наблюдений", "завершено_в_срок"])
        for moment, n, a, b in [("две недели", 400, 64, 54), ("месяц", 800, 125, 132), ("два месяца", 1600, 238, 284)]:
            writer.writerow([moment, "A", n, a])
            writer.writerow([moment, "B", n, b])
    print(f"created {len(rows)} training rows")


if __name__ == "__main__":
    main()
