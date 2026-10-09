#!/usr/bin/env python3
"""Validate downloadable workbooks and the seasonality story."""
from __future__ import annotations

import json
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]


def slope(values: list[float]) -> float:
    xbar = (len(values) - 1) / 2
    ybar = sum(values) / len(values)
    return sum((i - xbar) * (value - ybar) for i, value in enumerate(values)) / sum((i - xbar) ** 2 for i in range(len(values)))


seasonality = json.loads((ROOT / "data/government-seasonality.json").read_text(encoding="utf-8"))
assert len(seasonality["rows"]) == 2557
focus = seasonality["rows"][-seasonality["focus_days"] :]
signs = [slope([row[key] for row in focus]) > 0 for key in ("observed", "without_annual", "without_monthly", "baseline")]
assert signs == [True, False, True, False], signs

formula_path = ROOT / "downloads/uchebny_nabor_formuly_excel.xlsm"
formula_book = load_workbook(formula_path, read_only=True, keep_vba=True)
assert formula_book.sheetnames == ["Формулы", "Задание"]
assert formula_book["Формулы"].max_row == 121
with ZipFile(formula_path) as archive:
    assert "xl/vbaProject.bin" in archive.namelist()
    content_types = archive.read("[Content_Types].xml")
    assert b"application/vnd.ms-excel.sheet.macroEnabled.main+xml" in content_types

macro_path = ROOT / "downloads/obrashcheniya_12000_excel.xlsm"
macro_book = load_workbook(macro_path, read_only=True, keep_vba=True)
assert macro_book.sheetnames == ["Выгрузка", "Инструкция"]
assert macro_book["Выгрузка"].max_row == 12325
with ZipFile(macro_path) as archive:
    assert "xl/vbaProject.bin" in archive.namelist()
    content_types = archive.read("[Content_Types].xml")
    assert b"application/vnd.ms-excel.sheet.macroEnabled.main+xml" in content_types

assert not list((ROOT / "downloads").glob("*.xlsx")), "В публичных загрузках не должно оставаться XLSX"

print("training assets: passed")
