#!/usr/bin/env python3
"""Build deterministic public-sector training data and Excel workbooks."""
from __future__ import annotations

import csv
import json
import math
from datetime import date, timedelta
from pathlib import Path

import xlsxwriter
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DOWNLOADS = ROOT / "downloads"
DATA = ROOT / "data"


def build_seasonality() -> None:
    start = date(2019, 1, 1)
    days = 365 * 7 + 2
    focus = 35
    phase_week = 4.079158744652328
    phase_month = 4.9120680628743125
    phase_year = 5.737670202987268
    rows = []
    for i in range(days):
        local = i - (days - focus)
        baseline = 430 - 0.025 * i
        annual = 92.7345 * math.sin(2 * math.pi * local / 365.25 + phase_year)
        monthly = 58.9545 * math.sin(2 * math.pi * local / 30.44 + phase_month)
        weekly = 21.4815 * math.sin(2 * math.pi * local / 7 + phase_week)
        noise = 3.2 * math.sin(i * 1.71) + 1.7 * math.cos(i * 0.43)
        observed = baseline + annual + monthly + weekly + noise
        rows.append({
            "date": (start + timedelta(days=i)).isoformat(),
            "observed": round(observed, 2),
            "without_annual": round(observed - annual, 2),
            "without_monthly": round(observed - annual - monthly, 2),
            "baseline": round(observed - annual - monthly - weekly, 2),
            "annual": round(annual, 2),
            "monthly": round(monthly, 2),
            "weekly": round(weekly, 2),
        })
    payload = {
        "title": "Количество обращений, завершённых за день",
        "description": "Обезличенный обучающий ряд; значения не описывают работу реального органа власти.",
        "focus_days": focus,
        "rows": rows,
    }
    (DATA / "government-seasonality.json").write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    with (DOWNLOADS / "seasonality_government_training.csv").open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.DictWriter(file, fieldnames=rows[0].keys(), delimiter=";", lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def load_raw_rows() -> tuple[list[str], list[list[object]]]:
    source = DOWNLOADS / "obrashcheniya_12000_excel.xlsx"
    workbook = load_workbook(source, read_only=True, data_only=True)
    sheet = workbook["Выгрузка"]
    values = list(sheet.iter_rows(values_only=True))
    return list(values[0]), [list(row) for row in values[1:]]


def build_formula_workbook(headers: list[str], rows: list[list[object]]) -> None:
    target = DOWNLOADS / "uchebny_nabor_formuly_excel.xlsx"
    workbook = xlsxwriter.Workbook(target)
    sheet = workbook.add_worksheet("Формулы")
    guide = workbook.add_worksheet("Задание")
    title = workbook.add_format({"bold": True, "font_size": 16, "font_color": "#173A61"})
    head = workbook.add_format({"bold": True, "bg_color": "#173A61", "font_color": "#FFFFFF", "border": 1})
    date_fmt = workbook.add_format({"num_format": "dd.mm.yyyy", "border": 1})
    cell = workbook.add_format({"border": 1})
    guide.write("A1", "Практика: расчёты формулами Excel", title)
    guide.write("A3", "1. Рассчитайте длительность обработки: дата завершения − дата регистрации.")
    guide.write("A4", "2. Сравните длительность с целевым сроком и верните «Да» или «Нет».")
    guide.write("A5", "3. Посчитайте долю записей, завершённых в срок.")
    guide.write("A7", "Используйте только обезличенный обучающий набор. Лист «Формулы» содержит 120 строк с уже распознанными датами.")
    guide.set_column("A:A", 115)
    subset_headers = [headers[i] for i in (0, 1, 2, 3, 4, 5, 6)] + ["Длительность, дней", "Срок соблюдён"]
    for col, value in enumerate(subset_headers):
        sheet.write(0, col, value, head)
    for row_no, row in enumerate(rows[:120], 1):
        for target_col, source_col in enumerate((0, 1, 2, 3, 4, 5, 6)):
            value = row[source_col]
            if source_col in (1, 2) and hasattr(value, "year"):
                sheet.write_datetime(row_no, target_col, value, date_fmt)
            else:
                sheet.write(row_no, target_col, value, cell)
    sheet.add_table(0, 0, 120, len(subset_headers) - 1, {"name": "УчебныеДанные", "columns": [{"header": h} for h in subset_headers]})
    sheet.freeze_panes(1, 0)
    sheet.set_column(0, 0, 16)
    sheet.set_column(1, 2, 18)
    sheet.set_column(3, 8, 20)
    workbook.close()


def build_macro_workbook(headers: list[str], rows: list[list[object]]) -> None:
    target = DOWNLOADS / "obrashcheniya_12000_excel.xlsm"
    project = Path("/tmp/vbaProject.bin")
    if not project.exists():
        raise SystemExit("Expected /tmp/vbaProject.bin from the XlsxWriter example project")
    workbook = xlsxwriter.Workbook(target)
    workbook.add_vba_project(str(project))
    sheet = workbook.add_worksheet("Выгрузка")
    guide = workbook.add_worksheet("Инструкция")
    head = workbook.add_format({"bold": True, "bg_color": "#173A61", "font_color": "#FFFFFF", "border": 1})
    cell = workbook.add_format({"border": 1})
    title = workbook.add_format({"bold": True, "font_size": 16, "font_color": "#173A61"})
    guide.write("A1", "Продвинутая практика: обработка выгрузки с помощью VBA", title)
    guide.write("A3", "1. Получите модуль VBA в ИИ-сервисе и проверьте его в тренажёре.")
    guide.write("A4", "2. Нажмите Alt+F11, выберите Insert → Module и вставьте проверенный код.")
    guide.write("A5", "3. Вернитесь в Excel, нажмите Alt+F8 и запустите ProcessTrainingAppeals.")
    guide.write("A7", "Книга уже сохранена в формате XLSM: повторное сохранение в другой формат не требуется.")
    guide.write("A9", "Файл содержит обезличенные обучающие примеры и не описывает работу реальных организаций.")
    guide.set_column("A:A", 120)
    for col, value in enumerate(headers):
        sheet.write(0, col, value, head)
    for row_no, row in enumerate(rows, 1):
        for col, value in enumerate(row):
            sheet.write(row_no, col, value, cell)
    sheet.freeze_panes(1, 0)
    sheet.autofilter(0, 0, len(rows), len(headers) - 1)
    sheet.set_column(0, len(headers) - 1, 20)
    workbook.close()


if __name__ == "__main__":
    build_seasonality()
    source_headers, source_rows = load_raw_rows()
    build_formula_workbook(source_headers, source_rows)
    build_macro_workbook(source_headers, source_rows)
